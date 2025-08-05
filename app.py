

# import os
# import json
# import uuid
# from typing import List
# from fastapi import FastAPI, UploadFile, File, HTTPException
# from fastapi.middleware.cors import CORSMiddleware
# from pydantic import BaseModel
# from dotenv import load_dotenv
# import google.generativeai as genai
# from pinecone import Pinecone
# from upload_script import extract_text_from_pdf, semantic_chunk, upload_chunks_to_pinecone

# # === Load Env Vars ===
# load_dotenv()
# GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
# PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
# INDEX_NAME = os.getenv("PINECONE_INDEX_NAME")

# # === Initialize Clients ===
# genai.configure(api_key=GOOGLE_API_KEY)
# pc = Pinecone(api_key=PINECONE_API_KEY)
# index = pc.Index(INDEX_NAME)

# # === FastAPI App ===
# app = FastAPI(title="HackRx Policy LLM Backend")

# # CORS Middleware
# app.add_middleware(
#     CORSMiddleware,
#     allow_origins=["*"],
#     allow_methods=["*"],
#     allow_headers=["*"]
# )

# # === Request Schema ===
# class QueryRequest(BaseModel):
#     file_id: str
#     queries: List[str]  # ⬅️ Only an array of queries now

# # === Health Check ===
# @app.get("/health")
# async def health_check():
#     return {"status": "ok", "pinecone_index": INDEX_NAME}

# # === Upload PDF (Single PDF Upload) ===
# @app.post("/upload")
# async def upload_policy(file: UploadFile = File(...)):
#     try:
#         # Save temp file
#         file_path = f"./temp_{file.filename}"
#         with open(file_path, "wb") as f:
#             f.write(await file.read())

#         # Extract text & chunk
#         pdf_text = extract_text_from_pdf(file_path)
#         chunks = semantic_chunk(pdf_text)

#         # Generate unique file_id (namespace)
#         file_id = os.path.splitext(file.filename)[0] + "_" + str(uuid.uuid4())[:6]

#         # Upload chunks to Pinecone
#         upload_chunks_to_pinecone(chunks, file_id)

#         os.remove(file_path)  # Cleanup temp file

#         return {"message": "File uploaded and indexed successfully", "file_id": file_id}

#     except Exception as e:
#         raise HTTPException(status_code=500, detail=f"Upload failed: {e}")

# # === Query PDF (Array of Queries) ===
# @app.post("/query")
# async def query_policy(request: QueryRequest):
#     responses = []

#     for query_text in request.queries:
#         try:
#             # Embed Query
#             query_embedding = genai.embed_content(
#                 model="models/embedding-001",
#                 content=query_text,
#                 task_type="retrieval_query"
#             )["embedding"]

#             # Retrieve chunks from Pinecone (specific file namespace)
#             pinecone_results = index.query(
#                 vector=query_embedding,
#                 namespace=request.file_id,
#                 top_k=8,
#                 include_metadata=True
#             )

#             if not pinecone_results.get("matches"):
#                 responses.append({
#                     "decision": "Cannot determine",
#                     "amount": "Unknown",
#                     "justification": "No relevant info found.",
#                     "clause_ids": [],
#                     "risk_level": "Yellow",
#                     "answers": []
#                 })
#                 continue

#             context_chunks = [match["metadata"]["text"] for match in pinecone_results["matches"]]

#             # Build Gemini Prompt
#             prompt = f"""
#             You are an insurance policy assistant.
#             Use ONLY the clauses below to answer the query.

#             Clauses:
#             {'\n---\n'.join(context_chunks)}

#             Query:
#             {query_text}

#             Respond strictly in JSON:
#             {{
#                 "decision": "Yes / No / Cannot determine",
#                 "amount": "Coverage limit, percentage, or Unknown",
#                 "justification": "Explain using exact clauses",
#                 "clause_ids": ["..."],
#                 "risk_level": "Red / Orange / Yellow / Black",
#                 "answers": ["Direct factual answers"]
#             }}
#             """

#             model = genai.GenerativeModel("gemini-1.5-flash")
#             response = model.generate_content(prompt)
#             response_text = response.text.strip()

#             # Clean fenced JSON
#             if response_text.startswith("```json"):
#                 response_text = response_text[7:]
#             if response_text.endswith("```"):
#                 response_text = response_text[:-3]

#             try:
#                 responses.append(json.loads(response_text))
#             except:
#                 responses.append({
#                     "decision": "Cannot determine",
#                     "amount": "Unknown",
#                     "justification": "Parsing error.",
#                     "clause_ids": [],
#                     "risk_level": "Yellow",
#                     "answers": []
#                 })

#         except Exception as e:
#             responses.append({
#                 "decision": "Cannot determine",
#                 "amount": "Unknown",
#                 "justification": f"Error: {str(e)}",
#                 "clause_ids": [],
#                 "risk_level": "Yellow",
#                 "answers": []
#             })

#     return responses
       #raise HTTPException(status_code=500, detail=f"Webhook error: {str(e)}")

# import os
# import uuid
# import requests
# import tempfile
# from fastapi import FastAPI, HTTPException, Header
# from upload_script import extract_text_from_pdf, semantic_chunk, upload_chunks_to_pinecone
# from query_input import retrieve_chunks, ask_gemini

# app = FastAPI()
# @app.get("/")
# def root():
#     return {"status": "Server is running successfully!"}

# def download_pdf(pdf_input: str) -> str:
#     """Handles Google Drive, Direct PDF, and Base64 PDF download safely."""
#     tmp_dir = tempfile.gettempdir()
#     save_path = os.path.join(tmp_dir, "policy.pdf")

#     try:
#         # ✅ Case 1: Base64 PDF string
#         if pdf_input.startswith("data:application/pdf;base64,"):
#             import base64
#             pdf_data = base64.b64decode(pdf_input.split(",")[1])
#             with open(save_path, "wb") as f:
#                 f.write(pdf_data)
#             return save_path

#         # ✅ Case 2: Google Drive Links
#         if "drive.google.com" in pdf_input:
#             file_id = None
#             if "file/d/" in pdf_input:
#                 parts = pdf_input.split("/file/d/")
#                 if len(parts) > 1:
#                     file_id = parts[1].split("/")[0]
#             elif "id=" in pdf_input:
#                 parts = pdf_input.split("id=")
#                 if len(parts) > 1:
#                     file_id = parts[1].split("&")[0]

#             if not file_id:
#                 raise HTTPException(status_code=400, detail="Invalid Google Drive link format.")
            
#             pdf_url = f"https://drive.google.com/uc?export=download&id={file_id}"
#             session = requests.Session()
#             response = session.get(pdf_url, stream=True)

#             # Handle Google Drive virus scan confirmation pages
#             if "text/html" in response.headers.get("Content-Type", ""):
#                 for k, v in response.cookies.items():
#                     if k.startswith("download_warning"):
#                         confirm_url = f"{pdf_url}&confirm={v}"
#                         response = session.get(confirm_url, stream=True)
#                         break

#         else:
#             # ✅ Case 3: Direct PDF URLs
#             response = requests.get(pdf_input, stream=True, allow_redirects=True)

#         # ✅ Verify PDF signature (even if Content-Type header is missing)
#         content = response.content
#         if not content.startswith(b"%PDF"):
#             raise HTTPException(status_code=400, detail="Failed to download a valid PDF file.")

#         with open(save_path, "wb") as f:
#             f.write(content)

#         return save_path

#     except Exception as e:
#         raise HTTPException(status_code=400, detail=f"Failed to process PDF URL: {str(e)}")


# @app.post("/hackrx/run")
# async def hackrx_run(payload: dict, authorization: str = Header(None)):
#     try:
#         pdf_url = payload.get("documents")
#         questions = payload.get("questions")

#         if not pdf_url or not questions or not isinstance(questions, list) or len(questions) == 0:
#             raise HTTPException(status_code=400, detail="Missing or invalid 'documents' or 'questions'.")

#         # ✅ Step 1: Download PDF
#         pdf_path = download_pdf(pdf_url)

#         # ✅ Step 2: Extract text & Chunk
#         file_id = f"file_{uuid.uuid4().hex[:6]}"
#         pdf_text = extract_text_from_pdf(pdf_path)
#         if not pdf_text.strip():
#             raise HTTPException(status_code=400, detail="Extracted PDF text is empty.")
#         chunks = semantic_chunk(pdf_text)
#         upload_chunks_to_pinecone(chunks, file_id)
#         print("Chunks uploaded:", len(chunks))


#         # ✅ Step 3: Answer each question
#         answers = []
#         for q in questions:
#             try:
#                 retrieved_chunks = retrieve_chunks(q, file_id,top_k=8)
#                 print("\n--- Question:", q)
#                 print("Retrieved Chunks:\n", retrieved_chunks)

#                 result = ask_gemini(q, retrieved_chunks)
#                 answer_text = result["answers"][0] if result and "answers" in result and result["answers"] else "Cannot determine"
#                 answers.append(answer_text)
#             except Exception:
#                 answers.append("Cannot determine")

#         return {"answers": answers}

#     except Exception as e:
#         raise HTTPException(status_code=500, detail=f"Webhook error: {str(e)}")

# #nicely working code 
# import os
# import fitz  # PyMuPDF
# import uuid
# import re
# import hashlib
# import json
# import asyncio
# import requests
# from concurrent.futures import ThreadPoolExecutor
# from typing import List, Dict, Tuple
# from fastapi import FastAPI, UploadFile, File, Form, HTTPException
# from dotenv import load_dotenv
# from pinecone import Pinecone, ServerlessSpec
# from dataclasses import dataclass
# from pdf2image import convert_from_path
# import pytesseract
# import tempfile
# import time
# from rank_bm25 import BM25Okapi
# import google.generativeai as genai
# import email
# from email import policy
# from bs4 import BeautifulSoup
# from docx import Document

# # === Load Env Variables ===
# load_dotenv()

# GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
# PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
# INDEX_NAME = os.getenv("PINECONE_INDEX_NAME")
# PINECONE_CLOUD = os.getenv("PINECONE_CLOUD", "aws")
# PINECONE_REGION = os.getenv("PINECONE_REGION", "us-east-1")

# # === Configure Gemini & Pinecone ===
# genai.configure(api_key=GOOGLE_API_KEY)
# pc = Pinecone(api_key=PINECONE_API_KEY)

# if INDEX_NAME not in pc.list_indexes().names():
#     pc.create_index(
#         name=INDEX_NAME,
#         dimension=768,
#         metric="cosine",
#         spec=ServerlessSpec(cloud=PINECONE_CLOUD, region=PINECONE_REGION)
#     )
# index = pc.Index(INDEX_NAME)

# app = FastAPI()

# @dataclass
# class DocumentChunk:
#     text: str
#     chunk_id: str
#     page_num: int
#     chunk_type: str
#     metadata: Dict

# # === PDF Download (URL Support) ===
# def download_pdf(pdf_url: str) -> str:
#     tmp_path = os.path.join(tempfile.gettempdir(), "policy.pdf")
#     r = requests.get(pdf_url, stream=True)
#     if r.status_code != 200:
#         raise HTTPException(status_code=400, detail="Failed to download PDF.")
#     with open(tmp_path, "wb") as f:
#         f.write(r.content)
#     return tmp_path

# # === OCR Fallback ===
# def ocr_pdf(pdf_path: str) -> str:
#     pages = convert_from_path(pdf_path)
#     text = ""
#     for page in pages:
#         text += pytesseract.image_to_string(page)
#     return text

# # === PDF Text Extraction ===
# def extract_text_from_pdf_enhanced(pdf_path: str) -> Tuple[str, Dict]:
#     doc = fitz.open(pdf_path)
#     full_text = ""
#     structured_content = {'definitions': {}, 'clauses': {}, 'tables': [], 'key_terms': set()}
#     for page_num, page in enumerate(doc, start=1):
#         page_text = page.get_text("text").strip()
#         if page_text:
#             full_text += f"\n[Page {page_num}]\n{page_text}\n"
#             definition_matches = re.findall(r'([A-Z][a-z\s]+)\s*means\s+([^.]+\.)', page_text)
#             for term, definition in definition_matches:
#                 structured_content['definitions'][term.strip()] = definition.strip()
#                 structured_content['key_terms'].add(term.strip().lower())
#             clause_matches = re.findall(r'(\d+(?:\.\d+)*)\s+([^\n]+)', page_text)
#             for clause_num, clause_text in clause_matches:
#                 structured_content['clauses'][clause_num] = clause_text.strip()
#     return full_text, structured_content

# # === DOCX Extraction ===
# def extract_text_from_docx(file_path: str) -> Tuple[str, Dict]:
#     doc = Document(file_path)
#     full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
#     structured_content = {"headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]}
#     return full_text, structured_content

# # === Email Extraction ===
# def extract_text_from_email(file_path: str) -> Tuple[str, Dict]:
#     with open(file_path, "rb") as f:
#         msg = email.message_from_binary_file(f, policy=policy.default)

#     email_text = ""
#     structured_content = {"subject": msg["subject"], "from": msg["from"], "to": msg["to"], "body": ""}

#     for part in msg.walk():
#         if part.get_content_type() == "text/plain":
#             email_text += part.get_content()
#         elif part.get_content_type() == "text/html":
#             soup = BeautifulSoup(part.get_content(), "html.parser")
#             email_text += soup.get_text()

#     structured_content["body"] = email_text.strip()
#     return email_text, structured_content

# # === Smart Chunking ===
# def smart_semantic_chunk(text: str, structured_content: Dict, chunk_size: int = 400, overlap: int = 80) -> List[DocumentChunk]:
#     words = text.split()
#     chunks = []
#     for term, definition in structured_content.get('definitions', {}).items():
#         chunks.append(DocumentChunk(
#             text=f"{term} means {definition}",
#             chunk_id=f"def_{len(chunks)}",
#             page_num=0,
#             chunk_type='definition',
#             metadata={'term': term, 'importance': 'high'}
#         ))
#     for i in range(0, len(words), chunk_size - overlap):
#         chunk_text = " ".join(words[i:i + chunk_size])
#         if len(chunk_text.split()) >= 30:
#             chunk_type = 'general'
#             importance = 'medium'
#             if any(term in chunk_text.lower() for term in structured_content.get('key_terms', set())):
#                 chunk_type, importance = 'key_term', 'high'
#             elif re.search(r'\d+(?:\.\d+)*\s', chunk_text):
#                 chunk_type, importance = 'clause', 'high'
#             elif any(keyword in chunk_text.lower() for keyword in ['waiting period', 'coverage', 'premium', 'deductible']):
#                 chunk_type, importance = 'policy_term', 'high'
#             chunks.append(DocumentChunk(
#                 text=chunk_text.strip(),
#                 chunk_id=f"chunk_{len(chunks)}",
#                 page_num=i // 200 + 1,
#                 chunk_type=chunk_type,
#                 metadata={'importance': importance, 'word_count': len(chunk_text.split())}
#             ))
#     return chunks

# # === Embeddings (Gemini) ===
# def embed_text(text: str):
#     embedding = genai.embed_content(
#         model="models/embedding-001",
#         content=text,
#         task_type="retrieval_document"
#     )["embedding"]
#     return embedding

# def embed_text_for_query(text: str):
#     embedding = genai.embed_content(
#         model="models/embedding-001",
#         content=text,
#         task_type="retrieval_query"
#     )["embedding"]
#     return embedding

# # === Parallel Chunk Embedding ===
# async def embed_chunks_parallel(chunks: List[DocumentChunk], max_workers: int = 8) -> List[Tuple]:
#     def embed_single_chunk(chunk: DocumentChunk):
#         embedding = embed_text(chunk.text)
#         chunk_hash = hashlib.md5(chunk.text.encode()).hexdigest()
#         return (str(uuid.uuid4()), embedding, {
#             "text": chunk.text,
#             "hash": chunk_hash,
#             "chunk_type": chunk.chunk_type,
#             "importance": chunk.metadata.get('importance', 'medium'),
#             "page_num": chunk.page_num
#         })
#     with ThreadPoolExecutor(max_workers=max_workers) as executor:
#         return list(executor.map(embed_single_chunk, chunks))

# async def upload_chunks(chunks: List[DocumentChunk], file_id: str):
#     embedded = await embed_chunks_parallel(chunks)
#     index.upsert(vectors=embedded, namespace=file_id)

# # === Hybrid Search (BM25 + Vector) ===
# async def hybrid_search(query: str, file_id: str, bm25: BM25Okapi, chunks: List[str], top_k: int = 10):
#     async def bm25_search():
#         scores = bm25.get_scores(query.split())
#         top_indices = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:top_k]
#         return [chunks[i] for i in top_indices]

#     async def pinecone_search():
#         query_embedding = embed_text_for_query(query)
#         res = index.query(namespace=file_id, vector=query_embedding, top_k=top_k, include_metadata=True)
#         return [m["metadata"]["text"] for m in res["matches"]]

#     bm25_results, pinecone_results = await asyncio.gather(bm25_search(), pinecone_search())
#     seen, combined = set(), []
#     for c in bm25_results + pinecone_results:
#         if c not in seen:
#             seen.add(c)
#             combined.append(c)
#     return combined[:top_k]

# # === Batched Gemini QA ===import re
# import json
# import asyncio
# from fastapi import HTTPException

# async def ask_gemini_batch(questions: List[str], context_chunks: List[str], structured_content: Dict = None, retries: int = 3):
#     """
#     Ask Gemini model a batch of questions using provided context chunks.
#     Ensures detailed, clause-backed answers and avoids hallucination.
#     Includes robust JSON handling with fallback.
#     """

#     # Limit context chunks to avoid oversized prompt
#     context_chunks = context_chunks[:6]
#     combined_context = "\n\n--- CHUNK ---\n".join(context_chunks)
#     question_list = "\n".join([f"{i+1}. {q}" for i, q in enumerate(questions)])
    
#     definitions_section = ""
#     if structured_content and structured_content.get("definitions"):
#         definitions_section = f"""
# DEFINITIONS & KEY TERMS:
# {json.dumps(structured_content.get("definitions", {}), indent=2)}
# """

#     prompt = f"""
# You are an expert insurance analyst.
# Your job is to extract **accurate answers strictly from the document text provided**. 
# Do not infer or hallucinate any information not present in the text.

# DOCUMENT TEXT:
# {combined_context}

# {definitions_section}

# QUESTIONS:
# {question_list}

# Answer each question in **clear, complete sentences** with specific clause references or supporting context from the document where possible.
# If the answer is not explicitly stated in the text, respond with: "Not specified in the policy text."

# Output strictly in this JSON format:
# {{
#   "answers": [
#     "Answer 1 with full context and clause references if present",
#     "Answer 2 ...",
#     "Answer 3 ..."
#   ]
# }}
# """

#     for attempt in range(retries):
#         try:
#             model = genai.GenerativeModel("gemini-1.5-flash")
#             response = model.generate_content(
#                 prompt,
#                 generation_config={
#                     "response_mime_type": "application/json",
#                     "temperature": 0.0,
#                     "max_output_tokens": 1400
#                 }
#             )

#             raw_output = response.text.strip()
#             print(f"[DEBUG] Gemini Raw Output (Attempt {attempt+1}):\n{raw_output[:500]}")

#             # Try JSON parsing
#             try:
#                 parsed = json.loads(raw_output)
#             except json.JSONDecodeError:
#                 # Fallback: extract JSON substring if extra text is present
#                 match = re.search(r'\{[\s\S]*\}', raw_output)
#                 if match:
#                     parsed = json.loads(match.group(0))
#                 else:
#                     raise HTTPException(status_code=500, detail="Gemini returned invalid JSON.")

#             # Validate answers
#             if "answers" in parsed and isinstance(parsed["answers"], list):
#                 return parsed["answers"]
#             else:
#                 raise ValueError("Gemini returned JSON without 'answers' list.")

#         except Exception as e:
#             print(f"[ERROR] Gemini attempt {attempt+1} failed: {str(e)}")
#             if "429" in str(e) and attempt < retries - 1:
#                 await asyncio.sleep(2 ** attempt)  # exponential backoff
#             elif attempt == retries - 1:
#                 raise HTTPException(status_code=500, detail="Gemini QA failed after retries.")



# # === API Endpoint ===
# @app.post("/api/v1/hackrx/run")
# async def hackrx_run(file: UploadFile = File(None), pdf_url: str = Form(None), questions: str = Form(...)):
#     start_time = time.time()

#     try:
#         questions_list = json.loads(questions)
#     except:
#         raise HTTPException(status_code=400, detail="Invalid questions format. Must be a JSON list.")

#     if not file and not pdf_url:
#         raise HTTPException(status_code=400, detail="Provide either a file or pdf_url.")


#     # === File Handling ===
#     if pdf_url:
#         file_path = download_pdf(pdf_url)
#         file_type = "pdf"
#     else:
#         file_ext = file.filename.split(".")[-1].lower()
#         file_path = os.path.join(tempfile.gettempdir(), file.filename)
#         with open(file_path, "wb") as f:
#             f.write(await file.read())
#         file_type = file_ext

#     # === Extract Content Based on Type ===
#     if file_type == "pdf":
#         text, structured_content = extract_text_from_pdf_enhanced(file_path)
#         if not text.strip():
#             text = ocr_pdf(file_path)
#     elif file_type == "docx":
#         text, structured_content = extract_text_from_docx(file_path)
#     elif file_type in ["eml", "email"]:
#         text, structured_content = extract_text_from_email(file_path)
#     else:
#         raise HTTPException(status_code=400, detail="Unsupported file type. Only PDF, DOCX, and EML are allowed.")

#     # === Hash File & Handle Indexing ===
#     file_hash = hashlib.md5(text.encode()).hexdigest()
#     stats = index.describe_index_stats()

#     if file_hash not in stats.get("namespaces", {}):
#         chunks = smart_semantic_chunk(text, structured_content)
#         await upload_chunks(chunks, file_hash)
#         chunk_texts = [c.text for c in chunks]
#     else:
#         res = index.query(vector=[0.0]*768, namespace=file_hash, top_k=50, include_metadata=True)
#         chunk_texts = [m["metadata"]["text"] for m in res["matches"]]

#     # === Hybrid Search & QA ===
#     bm25 = BM25Okapi([chunk.split() for chunk in chunk_texts])
#     context_chunks = await hybrid_search(" ".join(questions_list), file_hash, bm25, chunk_texts)
#     answers = await ask_gemini_batch(questions_list, context_chunks,structured_content)

#     return {
#         "file_id": file_hash,
#         "file_type": file_type,
#         "answers": answers,
#         "processing_time": f"{time.time()-start_time:.2f}s"
#     }

# @app.get("/health")
# async def health_check():
#     return {"status": "healthy", "timestamp": time.time()}

# if __name__ == "__main__":
#     import uvicorn
#     uvicorn.run(app, host="0.0.0.0", port=8000)


# #well working 16/21
# import os
# import fitz  # PyMuPDF
# import uuid
# import re
# import hashlib
# import json
# import asyncio
# import requests
# from concurrent.futures import ThreadPoolExecutor
# from typing import List, Dict, Tuple
# from fastapi import FastAPI, UploadFile, File, Form, HTTPException
# from dotenv import load_dotenv
# from pinecone import Pinecone, ServerlessSpec
# from dataclasses import dataclass
# from pdf2image import convert_from_path
# import pytesseract
# import tempfile
# import time
# from rank_bm25 import BM25Okapi
# import google.generativeai as genai
# import email
# from email import policy
# from bs4 import BeautifulSoup
# from docx import Document

# # === Load Env Variables ===
# load_dotenv()
# GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
# PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
# INDEX_NAME = os.getenv("PINECONE_INDEX_NAME")
# PINECONE_CLOUD = os.getenv("PINECONE_CLOUD", "aws")
# PINECONE_REGION = os.getenv("PINECONE_REGION", "us-east-1")

# # === Configure Gemini & Pinecone ===
# genai.configure(api_key=GOOGLE_API_KEY)
# pc = Pinecone(api_key=PINECONE_API_KEY)
# if INDEX_NAME not in pc.list_indexes().names():
#     pc.create_index(name=INDEX_NAME, dimension=768, metric="cosine",
#                     spec=ServerlessSpec(cloud=PINECONE_CLOUD, region=PINECONE_REGION))
# index = pc.Index(INDEX_NAME)
# app = FastAPI()

# @dataclass
# class DocumentChunk:
#     text: str
#     chunk_id: str
#     page_num: int
#     chunk_type: str
#     metadata: Dict

# # === File Helpers ===
# def download_pdf(pdf_url: str) -> str:
#     tmp_path = os.path.join(tempfile.gettempdir(), "policy.pdf")
#     r = requests.get(pdf_url, stream=True)
#     if r.status_code != 200:
#         raise HTTPException(status_code=400, detail="Failed to download PDF.")
#     with open(tmp_path, "wb") as f:
#         f.write(r.content)
#     return tmp_path

# def ocr_pdf(pdf_path: str) -> str:
#     pages = convert_from_path(pdf_path)
#     return "".join([pytesseract.image_to_string(page) for page in pages])

# # === PDF Text Extraction ===
# def extract_text_from_pdf_enhanced(pdf_path: str) -> Tuple[str, Dict]:
#     doc = fitz.open(pdf_path)
#     full_text = ""
#     structured_content = {'definitions': {}, 'clauses': {}, 'tables': [], 'key_terms': set()}
#     for page_num, page in enumerate(doc, start=1):
#         page_text = page.get_text("text").strip()
#         if page_text:
#             full_text += f"\n[Page {page_num}]\n{page_text}\n"
#             # Extract definitions
#             for term, definition in re.findall(r'([A-Z][a-z\s]+)\s*means\s+([^.]+\.)', page_text):
#                 structured_content['definitions'][term.strip()] = definition.strip()
#                 structured_content['key_terms'].add(term.strip().lower())
#             # Extract clauses
#             for clause_num, clause_text in re.findall(r'(\d+(?:\.\d+)*)\s+([^\n]+)', page_text):
#                 structured_content['clauses'][clause_num] = clause_text.strip()
#     return full_text, structured_content

# def extract_text_from_docx(file_path: str) -> Tuple[str, Dict]:
#     doc = Document(file_path)
#     text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
#     structured_content = {"headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]}
#     return text, structured_content

# def extract_text_from_email(file_path: str) -> Tuple[str, Dict]:
#     with open(file_path, "rb") as f:
#         msg = email.message_from_binary_file(f, policy=policy.default)
#     email_text = ""
#     structured = {"subject": msg["subject"], "from": msg["from"], "to": msg["to"], "body": ""}
#     for part in msg.walk():
#         if part.get_content_type() == "text/plain":
#             email_text += part.get_content()
#         elif part.get_content_type() == "text/html":
#             email_text += BeautifulSoup(part.get_content(), "html.parser").get_text()
#     structured["body"] = email_text.strip()
#     return email_text, structured

# # === Enhanced Chunking ===
# def smart_semantic_chunk(text: str, structured_content: Dict, chunk_size=400, overlap=80) -> List[DocumentChunk]:
#     words = text.split()
#     chunks = []

#     # Definitions
#     for term, definition in structured_content.get('definitions', {}).items():
#         chunks.append(DocumentChunk(
#             text=f"Definition: {term} means {definition}",
#             chunk_id=f"def_{len(chunks)}",
#             page_num=0,
#             chunk_type='definition',
#             metadata={'importance': 'high'}
#         ))

#     # Clauses
#     for clause_num, clause_text in structured_content.get('clauses', {}).items():
#         chunks.append(DocumentChunk(
#             text=f"Clause {clause_num}: {clause_text}",
#             chunk_id=f"clause_{clause_num}",
#             page_num=0,
#             chunk_type='clause',
#             metadata={'importance': 'high'}
#         ))

#     # General content
#     for i in range(0, len(words), chunk_size - overlap):
#         chunk_text = " ".join(words[i:i+chunk_size])
#         if len(chunk_text.split()) >= 30:
#             chunks.append(DocumentChunk(
#                 text=chunk_text.strip(),
#                 chunk_id=f"chunk_{len(chunks)}",
#                 page_num=(i // 200) + 1,
#                 chunk_type='general',
#                 metadata={'importance': 'medium'}
#             ))
#     return chunks

# # === Embeddings ===
# def embed_text(text: str):
#     return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_document")["embedding"]

# def embed_text_for_query(text: str):
#     return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_query")["embedding"]

# async def embed_chunks_parallel(chunks: List[DocumentChunk]) -> List[Tuple]:
#     def embed_chunk(chunk: DocumentChunk):
#         emb = embed_text(chunk.text)
#         return (str(uuid.uuid4()), emb, {
#             "text": chunk.text,
#             "hash": hashlib.md5(chunk.text.encode()).hexdigest(),
#             "chunk_type": chunk.chunk_type,
#             "importance": chunk.metadata['importance'],
#             "page_num": chunk.page_num
#         })
#     with ThreadPoolExecutor(max_workers=8) as executor:
#         return list(executor.map(embed_chunk, chunks))

# async def upload_chunks(chunks: List[DocumentChunk], file_id: str):
#     vectors = await embed_chunks_parallel(chunks)
#     index.upsert(vectors=vectors, namespace=file_id)

# # === Hybrid Search ===
# async def hybrid_search(query: str, file_id: str, bm25: BM25Okapi, chunks: List[str], top_k=10):
#     async def bm25_search():
#         scores = bm25.get_scores(query.split())
#         top = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:top_k]
#         return [chunks[i] for i in top]

#     async def pinecone_search():
#         query_emb = embed_text_for_query(query)
#         res = index.query(namespace=file_id, vector=query_emb, top_k=top_k, include_metadata=True)
#         return [m["metadata"]["text"] for m in res["matches"]]

#     bm25_res, pinecone_res = await asyncio.gather(bm25_search(), pinecone_search())
#     seen, combined = set(), []
#     for c in bm25_res + pinecone_res:
#         if c not in seen:
#             seen.add(c)
#             combined.append(c)
#     return combined[:top_k]

# # === API Endpoint ===
# @app.post("/api/v1/hackrx/run")
# async def hackrx_run(file: UploadFile = File(None), pdf_url: str = Form(None), questions: str = Form(...)):
#     start_time = time.time()

#     try:
#         questions_list = json.loads(questions)
#     except:
#         raise HTTPException(status_code=400, detail="Invalid questions format. Must be a JSON list.")

#     if not file and not pdf_url:
#         raise HTTPException(status_code=400, detail="Provide either a file or pdf_url.")

#     # File Handling
#     if pdf_url:
#         file_path, file_type = download_pdf(pdf_url), "pdf"
#     else:
#         ext = file.filename.split(".")[-1].lower()
#         file_path = os.path.join(tempfile.gettempdir(), file.filename)
#         with open(file_path, "wb") as f:
#             f.write(await file.read())
#         file_type = ext

#     # Extraction
#     if file_type == "pdf":
#         text, structured = extract_text_from_pdf_enhanced(file_path)
#         if not text.strip(): text = ocr_pdf(file_path)
#     elif file_type == "docx": text, structured = extract_text_from_docx(file_path)
#     elif file_type in ["eml", "email"]: text, structured = extract_text_from_email(file_path)
#     else: raise HTTPException(status_code=400, detail="Unsupported file type.")

#     file_hash = hashlib.md5(text.encode()).hexdigest()
#     stats = index.describe_index_stats()

#     if file_hash not in stats.get("namespaces", {}):
#         chunks = smart_semantic_chunk(text, structured)
#         await upload_chunks(chunks, file_hash)
#         chunk_texts = [c.text for c in chunks]
#     else:
#         res = index.query(vector=[0.0]*768, namespace=file_hash, top_k=50, include_metadata=True)
#         chunk_texts = [m["metadata"]["text"] for m in res["matches"]]

#     bm25 = BM25Okapi([c.split() for c in chunk_texts])

#     # Per-question retrieval & Gemini QA
#     answers = []
#     for q in questions_list:
#         context_chunks = await hybrid_search(q, file_hash, bm25, chunk_texts, top_k=12)
#         prompt = f"""
#         You are an expert insurance analyst. Use only the text below.
        
#         DOCUMENT CONTEXT:
#         {chr(10).join(context_chunks)}

#         CLAUSES:
#         {json.dumps(structured.get('clauses', {}), indent=2)}

#         DEFINITIONS:
#         {json.dumps(structured.get('definitions', {}), indent=2)}

#         QUESTION:
#         {q}

#         Rules:
#         1. Quote exact policy text and clause references when possible.
#         2. If no relevant info is found, answer ONLY: "Not specified in the policy text."

#         Output JSON:
#         {{
#             "answers": ["<answer>"]
#         }}
#         """

#         model = genai.GenerativeModel("gemini-1.5-flash")
#         response = model.generate_content(prompt, generation_config={
#             "response_mime_type": "application/json",
#             "temperature": 0.0,
#             "max_output_tokens": 1000
#         })

#         try:
#             parsed = json.loads(response.text)
#             answer = parsed.get("answers", ["Not specified in the policy text."])[0]
#         except:
#             match = re.search(r'\{[\s\S]*\}', response.text)
#             answer = json.loads(match.group(0)).get("answers", ["Not specified in the policy text."])[0] if match else "Not specified in the policy text."

#         answers.append(answer)

#     return {"file_id": file_hash, "file_type": file_type, "answers": answers, "processing_time": f"{time.time()-start_time:.2f}s"}

# @app.get("/health")
# async def health_check():
#     return {"status": "healthy", "timestamp": time.time()}


#crazy code but processing time is 80sec+

# import os
# import fitz  # PyMuPDF
# import uuid
# import re
# import hashlib
# import json
# import asyncio
# import requests
# from concurrent.futures import ThreadPoolExecutor
# from typing import List, Dict, Tuple
# from fastapi import FastAPI, HTTPException, Body
# from dotenv import load_dotenv
# from pinecone import Pinecone, ServerlessSpec
# from dataclasses import dataclass
# from pdf2image import convert_from_path
# import pytesseract
# import tempfile
# import time
# from rank_bm25 import BM25Okapi
# import google.generativeai as genai
# import email
# from email import policy
# from bs4 import BeautifulSoup
# from docx import Document

# # === Load Env Variables ===
# load_dotenv()
# GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
# PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
# INDEX_NAME = os.getenv("PINECONE_INDEX_NAME")
# PINECONE_CLOUD = os.getenv("PINECONE_CLOUD", "aws")
# PINECONE_REGION = os.getenv("PINECONE_REGION", "us-east-1")

# # === Configure Gemini & Pinecone ===
# genai.configure(api_key=GOOGLE_API_KEY)
# pc = Pinecone(api_key=PINECONE_API_KEY)
# if INDEX_NAME not in pc.list_indexes().names():
#     pc.create_index(name=INDEX_NAME, dimension=768, metric="cosine",
#                     spec=ServerlessSpec(cloud=PINECONE_CLOUD, region=PINECONE_REGION))
# index = pc.Index(INDEX_NAME)
# app = FastAPI()


# @dataclass
# class DocumentChunk:
#     text: str
#     chunk_id: str
#     page_num: int
#     chunk_type: str
#     metadata: Dict


# # === File Helpers ===
# def download_file(file_url: str) -> Tuple[str, str]:
#     ext = file_url.split("?")[0].split(".")[-1].lower()
#     tmp_path = os.path.join(tempfile.gettempdir(), f"doc.{ext}")
#     r = requests.get(file_url, stream=True)
#     if r.status_code != 200:
#         raise HTTPException(status_code=400, detail="Failed to download file.")
#     with open(tmp_path, "wb") as f:
#         f.write(r.content)
#     return tmp_path, ext


# def ocr_pdf(pdf_path: str) -> str:
#     pages = convert_from_path(pdf_path)
#     return "".join([pytesseract.image_to_string(page) for page in pages])


# # === Extraction Functions ===
# def extract_text_from_pdf_enhanced(pdf_path: str) -> Tuple[str, Dict]:
#     doc = fitz.open(pdf_path)
#     full_text = ""
#     structured_content = {'definitions': {}, 'clauses': {}, 'tables': [], 'key_terms': set()}
#     for page_num, page in enumerate(doc, start=1):
#         page_text = page.get_text("text").strip()
#         if page_text:
#             full_text += f"\n[Page {page_num}]\n{page_text}\n"
#             # Definitions
#             for term, definition in re.findall(r'([A-Z][a-z\s]+)\s*means\s+([^.]+\.)', page_text):
#                 structured_content['definitions'][term.strip()] = definition.strip()
#                 structured_content['key_terms'].add(term.strip().lower())
#             # Clauses
#             for clause_num, clause_text in re.findall(r'(\d+(?:\.\d+)*)\s+([^\n]+)', page_text):
#                 structured_content['clauses'][clause_num] = clause_text.strip()
#     return full_text, structured_content


# def extract_text_from_docx(file_path: str) -> Tuple[str, Dict]:
#     doc = Document(file_path)
#     text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
#     structured_content = {"headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]}
#     return text, structured_content


# def extract_text_from_email(file_path: str) -> Tuple[str, Dict]:
#     with open(file_path, "rb") as f:
#         msg = email.message_from_binary_file(f, policy=policy.default)
#     email_text = ""
#     structured = {"subject": msg["subject"], "from": msg["from"], "to": msg["to"], "body": ""}
#     for part in msg.walk():
#         if part.get_content_type() == "text/plain":
#             email_text += part.get_content()
#         elif part.get_content_type() == "text/html":
#             email_text += BeautifulSoup(part.get_content(), "html.parser").get_text()
#     structured["body"] = email_text.strip()
#     return email_text, structured


# # === Chunking ===
# def smart_semantic_chunk(text: str, structured_content: Dict, chunk_size=400, overlap=80) -> List[DocumentChunk]:
#     words = text.split()
#     chunks = []

#     # Definitions
#     for term, definition in structured_content.get('definitions', {}).items():
#         chunks.append(DocumentChunk(
#             text=f"Definition: {term} means {definition}",
#             chunk_id=f"def_{len(chunks)}",
#             page_num=0,
#             chunk_type='definition',
#             metadata={'importance': 'high'}
#         ))

#     # Clauses
#     for clause_num, clause_text in structured_content.get('clauses', {}).items():
#         chunks.append(DocumentChunk(
#             text=f"Clause {clause_num}: {clause_text}",
#             chunk_id=f"clause_{clause_num}",
#             page_num=0,
#             chunk_type='clause',
#             metadata={'importance': 'high'}
#         ))

#     # General content
#     for i in range(0, len(words), chunk_size - overlap):
#         chunk_text = " ".join(words[i:i + chunk_size])
#         if len(chunk_text.split()) >= 30:
#             chunks.append(DocumentChunk(
#                 text=chunk_text.strip(),
#                 chunk_id=f"chunk_{len(chunks)}",
#                 page_num=(i // 200) + 1,
#                 chunk_type='general',
#                 metadata={'importance': 'medium'}
#             ))
#     return chunks


# # === Embeddings ===
# def embed_text(text: str):
#     return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_document")["embedding"]


# def embed_text_for_query(text: str):
#     return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_query")["embedding"]


# async def embed_chunks_parallel(chunks: List[DocumentChunk]) -> List[Tuple]:
#     def embed_chunk(chunk: DocumentChunk):
#         emb = embed_text(chunk.text)
#         return (str(uuid.uuid4()), emb, {
#             "text": chunk.text,
#             "hash": hashlib.md5(chunk.text.encode()).hexdigest(),
#             "chunk_type": chunk.chunk_type,
#             "importance": chunk.metadata['importance'],
#             "page_num": chunk.page_num
#         })

#     with ThreadPoolExecutor(max_workers=8) as executor:
#         return list(executor.map(embed_chunk, chunks))


# async def upload_chunks(chunks: List[DocumentChunk], file_id: str):
#     vectors = await embed_chunks_parallel(chunks)
#     index.upsert(vectors=vectors, namespace=file_id)


# # === Hybrid Search ===
# async def hybrid_search(query: str, file_id: str, bm25: BM25Okapi, chunks: List[str], top_k=10):
#     async def bm25_search():
#         scores = bm25.get_scores(query.split())
#         top = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:top_k]
#         return [chunks[i] for i in top]

#     async def pinecone_search():
#         query_emb = embed_text_for_query(query)
#         res = index.query(namespace=file_id, vector=query_emb, top_k=top_k, include_metadata=True)
#         return [m["metadata"]["text"] for m in res["matches"]]

#     bm25_res, pinecone_res = await asyncio.gather(bm25_search(), pinecone_search())
#     seen, combined = set(), []
#     for c in bm25_res + pinecone_res:
#         if c not in seen:
#             seen.add(c)
#             combined.append(c)
#     return combined[:top_k]


# # === JSON Sanitizer ===
# def extract_valid_json(text: str) -> dict:
#     match = re.search(r'\{[\s\S]*\}', text)
#     if not match:
#         return {"answers": ["Not specified in the policy text."]}
#     json_str = match.group(0)
#     json_str = re.sub(r",\s*}", "}", json_str)
#     json_str = re.sub(r",\s*]", "]", json_str)
#     try:
#         return json.loads(json_str)
#     except:
#         return {"answers": ["Not specified in the policy text."]}


# # === API Endpoint ===
# @app.post("/api/v1/hackrx/run")
# async def hackrx_run(payload: dict = Body(...)):
#     start_time = time.time()

#     pdf_url = payload.get("documents")
#     questions_list = payload.get("questions")

#     if not pdf_url or not isinstance(questions_list, list):
#         raise HTTPException(status_code=400, detail="Invalid payload. Provide 'documents' (URL) and 'questions' (list).")

#     file_path, file_type = download_file(pdf_url)

#     if file_type == "pdf":
#         text, structured = extract_text_from_pdf_enhanced(file_path)
#         if not text.strip(): text = ocr_pdf(file_path)
#     elif file_type == "docx":
#         text, structured = extract_text_from_docx(file_path)
#     elif file_type in ["eml", "email"]:
#         text, structured = extract_text_from_email(file_path)
#     else:
#         raise HTTPException(status_code=400, detail="Unsupported file type.")

#     file_hash = hashlib.md5(text.encode()).hexdigest()
#     stats = index.describe_index_stats()

#     if file_hash not in stats.get("namespaces", {}):
#         chunks = smart_semantic_chunk(text, structured)
#         await upload_chunks(chunks, file_hash)
#         chunk_texts = [c.text for c in chunks]
#     else:
#         res = index.query(vector=[0.0] * 768, namespace=file_hash, top_k=50, include_metadata=True)
#         chunk_texts = [m["metadata"]["text"] for m in res["matches"]]

#     bm25 = BM25Okapi([c.split() for c in chunk_texts])

#     async def process_question(q):
#         context_chunks = await hybrid_search(q, file_hash, bm25, chunk_texts, top_k=12)
#         prompt = f"""
#         You are an expert insurance analyst. Use only the text below.

#         DOCUMENT CONTEXT:
#         {chr(10).join(context_chunks)}

#         CLAUSES:
#         {json.dumps(structured.get('clauses', {}), indent=2)}

#         DEFINITIONS:
#         {json.dumps(structured.get('definitions', {}), indent=2)}

#         QUESTION:
#         {q}

#         Rules:
#         1. Quote exact policy text and clause references when possible.
#         2. If no relevant info is found, answer ONLY: "Not specified in the policy text."

#         Output JSON:
#         {{
#             "answers": ["<answer>"]
#         }}
#         """

#         model = genai.GenerativeModel("gemini-1.5-flash")
#         response = model.generate_content(
#             prompt,
#             generation_config={
#                 "response_mime_type": "application/json",
#                 "response_schema": {
#                     "type": "object",
#                     "properties": {
#                         "answers": {
#                             "type": "array",
#                             "items": {"type": "string"}
#                         }
#                     },
#                     "required": ["answers"]
#                 },
#                 "temperature": 0.0,
#                 "max_output_tokens": 1200
#             }
#         )
#         parsed = extract_valid_json(response.text)
#         return parsed.get("answers", ["Not specified in the policy text."])[0]

#     answers = await asyncio.gather(*(process_question(q) for q in questions_list))

#     return {
#         "file_id": file_hash,
#         "file_type": file_type,
#         "answers": answers,
#         "processing_time": f"{time.time()-start_time:.2f}s"
#     }


# @app.get("/health")
# async def health_check():
#     return {"status": "healthy", "timestamp": time.time()}


import os
import fitz
import uuid
import re
import hashlib
import json
import asyncio
import requests
import tempfile
import time
from typing import List, Dict, Tuple
from fastapi import FastAPI, HTTPException, Body
from dotenv import load_dotenv
from pinecone import Pinecone, ServerlessSpec
from dataclasses import dataclass
from pdf2image import convert_from_path
import pytesseract
from rank_bm25 import BM25Okapi
import google.generativeai as genai
from bs4 import BeautifulSoup
from docx import Document
import email
from email import policy

# === Load Env ===
load_dotenv()
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
INDEX_NAME = os.getenv("PINECONE_INDEX_NAME")
PINECONE_CLOUD = os.getenv("PINECONE_CLOUD", "aws")
PINECONE_REGION = os.getenv("PINECONE_REGION", "us-east-1")

# === Configure Gemini & Pinecone ===
genai.configure(api_key=GOOGLE_API_KEY)
pc = Pinecone(api_key=PINECONE_API_KEY)
if INDEX_NAME not in pc.list_indexes().names():
    pc.create_index(name=INDEX_NAME, dimension=768, metric="cosine",
                    spec=ServerlessSpec(cloud=PINECONE_CLOUD, region=PINECONE_REGION))
index = pc.Index(INDEX_NAME)

app = FastAPI()

@dataclass
class DocumentChunk:
    text: str
    chunk_id: str
    page_num: int
    chunk_type: str
    metadata: Dict

# === File Helpers ===
def download_file(url: str) -> str:
    tmp_path = os.path.join(tempfile.gettempdir(), "doc_download")
    os.makedirs(tmp_path, exist_ok=True)
    local_path = os.path.join(tmp_path, os.path.basename(url.split("?")[0]))
    r = requests.get(url, stream=True, timeout=30)
    if r.status_code != 200:
        raise HTTPException(status_code=400, detail="Failed to download file.")
    with open(local_path, "wb") as f:
        f.write(r.content)
    return local_path

def ocr_pdf(pdf_path: str) -> str:
    pages = convert_from_path(pdf_path)
    return "".join([pytesseract.image_to_string(page) for page in pages])

def extract_pdf(pdf_path: str) -> Tuple[str, Dict]:
    doc = fitz.open(pdf_path)
    text = ""
    structured = {"definitions": {}, "clauses": {}, "tables": [], "key_terms": set()}
    for i, page in enumerate(doc, start=1):
        page_text = page.get_text("text").strip()
        if page_text:
            text += f"\n[Page {i}]\n{page_text}\n"
            for term, definition in re.findall(r'([A-Z][a-z\s]+)\s*means\s+([^.]+\.)', page_text):
                structured['definitions'][term.strip()] = definition.strip()
            for clause_num, clause_text in re.findall(r'(\d+(?:\.\d+)*)\s+([^\n]+)', page_text):
                structured['clauses'][clause_num] = clause_text.strip()
    return text, structured

def extract_docx(file_path: str) -> Tuple[str, Dict]:
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    structured = {"headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]}
    return text, structured

def extract_email(file_path: str) -> Tuple[str, Dict]:
    with open(file_path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    email_text = ""
    structured = {"subject": msg["subject"], "from": msg["from"], "to": msg["to"], "body": ""}
    for part in msg.walk():
        if part.get_content_type() == "text/plain":
            email_text += part.get_content()
        elif part.get_content_type() == "text/html":
            email_text += BeautifulSoup(part.get_content(), "html.parser").get_text()
    structured["body"] = email_text.strip()
    return email_text, structured

# === Chunking ===
def smart_chunk(text: str, structured: Dict, chunk_size=400, overlap=80) -> List[DocumentChunk]:
    words = text.split()
    chunks = []
    for term, definition in structured.get('definitions', {}).items():
        chunks.append(DocumentChunk(f"Definition: {term} means {definition}", f"def_{len(chunks)}", 0, "definition", {}))
    for clause, c_text in structured.get('clauses', {}).items():
        chunks.append(DocumentChunk(f"Clause {clause}: {c_text}", f"clause_{clause}", 0, "clause", {}))
    for i in range(0, len(words), chunk_size - overlap):
        c_text = " ".join(words[i:i+chunk_size])
        if len(c_text.split()) >= 25:
            chunks.append(DocumentChunk(c_text, f"chunk_{len(chunks)}", (i // 200) + 1, "general", {}))
    return chunks

# === Embedding & Search ===
def embed_text(text: str):
    return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_document")["embedding"]

def embed_query(text: str):
    return genai.embed_content(model="models/embedding-001", content=text, task_type="retrieval_query")["embedding"]

async def upload_chunks(chunks: List[DocumentChunk], file_id: str):
    loop = asyncio.get_event_loop()
    vectors = await asyncio.gather(*[
        loop.run_in_executor(None, lambda c=chunk: (str(uuid.uuid4()), embed_text(c.text),
        {"text": c.text, "chunk_type": c.chunk_type, "page_num": c.page_num}))
        for chunk in chunks
    ])
    index.upsert(vectors=vectors, namespace=file_id)

async def hybrid_search(query: str, file_id: str, bm25: BM25Okapi, chunks: List[str], top_k=8):
    async def bm25_top():
        scores = bm25.get_scores(query.split())
        top_idx = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:top_k]
        return [chunks[i] for i in top_idx]
    async def pinecone_top():
        res = index.query(namespace=file_id, vector=embed_query(query), top_k=top_k, include_metadata=True)
        return [m["metadata"]["text"] for m in res["matches"]]
    bm25_res, pinecone_res = await asyncio.gather(bm25_top(), pinecone_top())
    return list(dict.fromkeys(bm25_res + pinecone_res))[:top_k]

# === Batched Gemini Answering ===
async def batch_gemini_answer(questions: List[str], context: str) -> List[str]:
    q_json = json.dumps(questions, indent=2)
    prompt = f"""
You are an insurance policy expert. Use ONLY the context below to answer.

DOCUMENT CONTEXT:
{context}

QUESTIONS:
{q_json}

Rules:
1. Answer each question accurately.
2. If answer is missing, respond "Not specified in the policy text."
Output JSON strictly: {{"answers": [<ans1>, <ans2>, ...]}}
"""
    model = genai.GenerativeModel("gemini-1.5-flash")
    resp = model.generate_content(prompt, generation_config={"temperature": 0, "response_mime_type": "application/json"})
    try:
        return json.loads(resp.text)["answers"]
    except:
        match = re.search(r'\{[\s\S]*\}', resp.text)
        return json.loads(match.group(0)).get("answers", ["Not specified in the policy text."]*len(questions))

# === Main Endpoint ===
@app.post("/api/v1/hackrx/run")
async def hackrx_run(payload: dict = Body(...)):
    start = time.time()
    pdf_url = payload.get("documents")
    questions = payload.get("questions")
    if not pdf_url or not questions:
        raise HTTPException(400, "Missing 'documents' or 'questions'")
    
    file_path = download_file(pdf_url)
    ext = file_path.split(".")[-1].lower()

    # Extract text
    if ext == "pdf":
        text, structured = extract_pdf(file_path)
        if not text.strip(): text = ocr_pdf(file_path)
    elif ext == "docx":
        text, structured = extract_docx(file_path)
    elif ext in ["eml", "email"]:
        text, structured = extract_email(file_path)
    else:
        raise HTTPException(400, "Unsupported file format.")

    file_hash = hashlib.md5(text.encode()).hexdigest()
    stats = index.describe_index_stats()
    if file_hash not in stats.get("namespaces", {}):
        chunks = smart_chunk(text, structured)
        await upload_chunks(chunks, file_hash)
        chunk_texts = [c.text for c in chunks]
    else:
        res = index.query(vector=[0.0]*768, namespace=file_hash, top_k=50, include_metadata=True)
        chunk_texts = [m["metadata"]["text"] for m in res["matches"]]

    bm25 = BM25Okapi([c.split() for c in chunk_texts])

    # Retrieve and batch Gemini answers
    answers = []
    batch_size = 4
    for i in range(0, len(questions), batch_size):
        batch_qs = questions[i:i+batch_size]
        contexts = await asyncio.gather(*[hybrid_search(q, file_hash, bm25, chunk_texts) for q in batch_qs])
        merged_context = "\n\n".join(["\n".join(c) for c in contexts])
        answers.extend(await batch_gemini_answer(batch_qs, merged_context))

    return {
        "file_id": file_hash,
        "file_type": ext,
        "answers": answers,
        "processing_time": f"{time.time()-start:.2f}s"
    }
